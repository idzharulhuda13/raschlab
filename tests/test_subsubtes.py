"""SUBSUBTES column and per-sub-subtes summary (2026-09-30).

Real TBS label mapping, the frozen item-table tail (ITEM, SUBSUBTES, STATUS),
the subsubtes_summary.csv shape (header-only when nothing maps) and the
'subsubtes' sheet in analysis_report.xlsx.
"""
import csv

import numpy as np

from raschlab.cli import run_analyze
from raschlab.report import (
    ITEM_HEADER_ROW_1,
    ITEM_HEADER_ROW_2,
    SUBSUBTES_SUMMARY_COLUMNS,
    item_table_rows,
    subsubtes_name,
    subsubtes_summary_rows,
)

FIT = {
    "se": np.array([0.5, 0.5, 0.5]),
    "infit_mnsq": np.array([1.0, 1.6, 0.9]),
    "infit_zstd": np.array([0.1, 2.0, -0.2]),
    "outfit_mnsq": np.array([1.0, 1.1, 0.8]),
    "outfit_zstd": np.array([0.0, 1.0, -0.5]),
}
X = np.array([[1.0, 0.0, 1.0], [0.0, 1.0, 1.0], [1.0, 1.0, 0.0]])
MASK = np.ones_like(X, dtype=bool)
LABELS = ["06tbspl25a06", "07tbspl25a07", "01tbskab26a01"]


def _rows(labels=LABELS):
    return item_table_rows(
        X, MASK, "AAA", np.array([1.2, -0.4, 0.0]), FIT, item_labels=labels
    )


def test_mapping_real_and_unknown_labels():
    assert subsubtes_name("06tbspl25a06") == "Logis"
    assert subsubtes_name("01tbskab26a01") == "Aritmatika dan Aljabar"
    assert subsubtes_name("06TBSPL25A06") == "Logis"
    assert subsubtes_name("00xyz00") == ""
    assert subsubtes_name("") == ""
    assert subsubtes_name(None) == ""


def test_item_row_tail_is_ITEM_SUBSUBTES_STATUS():
    rows = _rows()
    assert len(ITEM_HEADER_ROW_1) == len(ITEM_HEADER_ROW_2) == 16
    assert ITEM_HEADER_ROW_2[-3:] == ["ITEM", "SUBSUBTES", "STATUS"]
    assert list(rows[0].keys())[-3:] == ["ITEM", "SUBSUBTES", "STATUS"]
    assert [r["SUBSUBTES"] for r in rows] == [
        "Logis", "Logis", "Aritmatika dan Aljabar",
    ]


def test_label_less_rows_carry_empty_subsubtes():
    rows = _rows(labels=None)
    assert all(r["ITEM"] == "" for r in rows)
    assert all(r["SUBSUBTES"] == "" for r in rows)
    assert all(r["STATUS"] == "kept" for r in rows)


def test_summary_rows_shape_order_and_misfit_threshold():
    assert SUBSUBTES_SUMMARY_COLUMNS == [
        "SUBSUBTES", "ITEMS", "ANCHOR_ITEMS", "NEW_ITEMS",
        "MEAN_MEASURE", "S.SD_MEASURE", "MEAN_INFIT", "MAX_INFIT", "MISFIT_ITEMS",
    ]
    out = subsubtes_summary_rows(_rows())
    assert [r[0] for r in out] == ["Logis", "Aritmatika dan Aljabar"]
    assert out[0][1:4] == [2, 2, 0]
    assert out[1][1:4] == [1, 0, 1]
    assert out[0][4] == 0.4
    assert out[0][7] == 1.6
    assert out[0][8] == 1
    assert out[1][8] == 0
    assert subsubtes_summary_rows(_rows(labels=["zzz", "yyy", "qqq"])) == []


def test_end_to_end_writes_summary_csv_and_sheet(tmp_path):
    con = tmp_path / "c.con"
    con.write_text(
        "&INST\nITEM1 = 6\nNI = 3\nNAME1 = 1\nNAMLEN = 4\nKEY1 = ABC\n&END\n",
        encoding="utf-8",
    )
    prn = tmp_path / "d.prn"
    prn.write_text(
        "P001 ABC\nP002 BCA\nP003 CAB\nP004 AAB\nP005 BBA\nP006 CCA\n",
        encoding="utf-8",
    )
    lab = tmp_path / "labels.txt"
    lab.write_text("06tbspl25a06\n06tbspl25a07\n01tbskab26a01\n", encoding="utf-8")
    out = tmp_path / "out"
    run_analyze(str(con), str(prn), labels_path=str(lab), out_dir=str(out), out_format="both")

    with open(out / "subsubtes_summary.csv", newline="", encoding="utf-8") as f:
        data = list(csv.reader(f))
    assert data[0] == SUBSUBTES_SUMMARY_COLUMNS
    assert [r[0] for r in data[1:]] == ["Logis", "Aritmatika dan Aljabar"]

    import openpyxl
    wb = openpyxl.load_workbook(out / "analysis_report.xlsx")
    assert "subsubtes" in wb.sheetnames
    assert "tabulasi" in wb.sheetnames
    assert wb.sheetnames[-1] == "tabulasi"

    with open(out / "item_table_15.1.csv", newline="", encoding="utf-8") as f:
        item = list(csv.reader(f))
    assert item[1][-1] == "STATUS" and item[1][-2] == "SUBSUBTES"
    assert [r[14] for r in item[2:]] == ["Logis", "Logis", "Aritmatika dan Aljabar"]
