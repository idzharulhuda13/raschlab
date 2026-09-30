"""Tests for the item-delete (IDFILE) contract.

Fixture: /tmp/raschlab_demo/{demo.CON, demo_data.prn, idfile.TXT}
idfile.TXT contains a single entry: 6  (delete item 6 out of NI=6)
"""

import csv
import io
import os
import shutil
import sys
from unittest.mock import patch

import openpyxl
import pytest

from raschlab.cli import run_analyze

DEMO_DIR = "/tmp/raschlab_demo"
DEMO_CON = os.path.join(DEMO_DIR, "demo.CON")
DEMO_DATA = os.path.join(DEMO_DIR, "demo_data.prn")
DEMO_IDFILE = os.path.join(DEMO_DIR, "idfile.TXT")


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _run(tmp_path, *, idfile_path=None, con_path=None, capture_stderr=False):
    """Run run_analyze on the demo fixture into tmp_path.

    Returns (stdout_str, stderr_str, exit_code).
    """
    out_dir = str(tmp_path)
    buf_out = io.StringIO()
    buf_err = io.StringIO()
    code = 0
    with patch("sys.stdout", buf_out), patch("sys.stderr", buf_err):
        try:
            run_analyze(
                con_path=con_path or DEMO_CON,
                data_path=DEMO_DATA,
                idfile_path=idfile_path,
                out_dir=out_dir,
                out_format="both",
            )
        except SystemExit as e:
            code = e.code
    return buf_out.getvalue(), buf_err.getvalue(), code


def _read_csv(path):
    """Read CSV and return (header_rows, data_rows) skipping the 2-row header block."""
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    return rows[:2], rows[2:]


def _item_rows(tmp_path):
    _, data = _read_csv(os.path.join(str(tmp_path), "item_table_15.1.csv"))
    return data


def _summary_dict(tmp_path):
    """Return {(section, stat): value} from summary_table.csv."""
    _, rows = _read_csv(os.path.join(str(tmp_path), "summary_table.csv"))
    return {(r[0].strip(), r[1].strip()): r[2].strip() for r in rows if len(r) >= 3}


# ---------------------------------------------------------------------------
# 1. IDFILE= control-file resolution AND --idfile flag
# ---------------------------------------------------------------------------

def test_idfile_flag_resolves(tmp_path):
    """--idfile flag: item 6 ends up deleted."""
    out, err, code = _run(tmp_path, idfile_path=DEMO_IDFILE)
    assert code == 0, err
    rows = _item_rows(tmp_path)
    # 2026-09-30: SUBSUBTES occupies index 14 in the item table, so STATUS moved
    # from index 14 to 15.
    statuses = [r[15].strip() for r in rows]
    assert statuses.count("deleted") == 1
    assert statuses.count("kept") == 5


def test_idfile_con_basename_resolution(tmp_path):
    """IDFILE= basename in .CON next to itself is resolved without --idfile flag."""
    # Build a CON that names IDFILE= by basename only
    con_dir = tmp_path / "con_dir"
    con_dir.mkdir()
    con_path = con_dir / "test.CON"
    idfile_path = con_dir / "myidfile.TXT"
    idfile_path.write_text("6\n", encoding="utf-8")

    con_content = (
        "&INST\n"
        "ITEM1 = 10\n"
        "NI = 6\n"
        "NAME1 = 1\n"
        "NAMLEN = 8\n"
        "KEY1 = ABCDEF\n"
        "CODES = ABCDEF\n"
        "MISSCORE = -1\n"
        "IDFILE = myidfile.TXT\n"
        "&END\n"
    )
    con_path.write_text(con_content, encoding="utf-8")

    out_dir = tmp_path / "out"
    out, err, code = _run(out_dir, idfile_path=None, con_path=str(con_path))
    assert code == 0, err
    rows = _item_rows(out_dir)
    assert any(r[15].strip() == "deleted" for r in rows)


# ---------------------------------------------------------------------------
# 2. Out-of-range -> exit 2 with range message; malformed -> exit 2 "Error reading IDFILE"
# ---------------------------------------------------------------------------

def test_idfile_out_of_range_exits_2(tmp_path):
    idf = tmp_path / "oor.TXT"
    idf.write_text("99\n", encoding="utf-8")  # item 99 > NI=6
    out, err, code = _run(tmp_path / "out", idfile_path=str(idf))
    assert code == 2
    assert "outside 1..6" in err


def test_idfile_malformed_exits_2(tmp_path):
    idf = tmp_path / "bad.TXT"
    idf.write_text("notanumber\n", encoding="utf-8")
    out, err, code = _run(tmp_path / "out", idfile_path=str(idf))
    assert code == 2
    assert "Error reading IDFILE" in err


# ---------------------------------------------------------------------------
# 3. Empty idfile -> no deletes; item_table_15.1.csv identical to run without idfile
# ---------------------------------------------------------------------------

def test_empty_idfile_matches_baseline(tmp_path):
    base_dir = tmp_path / "base"
    idf_dir = tmp_path / "idf"
    idf = tmp_path / "empty.TXT"
    idf.write_text("", encoding="utf-8")

    _run(base_dir)
    _run(idf_dir, idfile_path=str(idf))

    base_content = (base_dir / "item_table_15.1.csv").read_bytes()
    idf_content = (idf_dir / "item_table_15.1.csv").read_bytes()
    assert base_content == idf_content


# ---------------------------------------------------------------------------
# 4. Deleted item leaves calibration: person SCORE/COUNT recomputed on reduced set
#    Deleting item 6 flips demo person 3 to 5/5 (extreme_max).
#    Console: "Extreme count 5 (1 min, 4 max)"
# ---------------------------------------------------------------------------

def test_delete_item6_person3_becomes_extreme(tmp_path):
    out, err, code = _run(tmp_path, idfile_path=DEMO_IDFILE)
    assert code == 0, err

    # Baseline: person 3 (DEMO0003) has 6 responses and score=5 (not all-correct over 5 items)
    # After item 6 deleted: DEMO0003 scores 5/5 -> extreme_max
    _, rows = _read_csv(tmp_path / "person_table.csv")
    person3 = next(r for r in rows if r[13].strip() == "DEMO0003")
    assert person3[1].strip() == "5", "DEMO0003 score should be 5"
    assert person3[2].strip() == "5", "DEMO0003 count should be 5"
    assert person3[17].strip() == "extreme_max", "DEMO0003 should be extreme_max"


def test_delete_item6_console_extreme_count(tmp_path):
    out, err, code = _run(tmp_path, idfile_path=DEMO_IDFILE)
    assert code == 0, err
    assert "Extreme count      : 5 (1 min, 4 max)" in out


# ---------------------------------------------------------------------------
# 5. Kept item rows: SCORE and COUNT unchanged vs baseline (pristine matrix)
# ---------------------------------------------------------------------------

def test_kept_items_score_count_unchanged(tmp_path):
    base_dir = tmp_path / "base"
    del_dir = tmp_path / "del"
    _run(base_dir)
    _run(del_dir, idfile_path=DEMO_IDFILE)

    base_rows = {r[0].strip(): r for r in _item_rows(base_dir)}
    del_rows = {r[0].strip(): r for r in _item_rows(del_dir)}

    for entry in ("1", "2", "3", "4", "5"):  # kept items
        assert del_rows[entry][1].strip() == base_rows[entry][1].strip(), f"item {entry} SCORE mismatch"
        assert del_rows[entry][2].strip() == base_rows[entry][2].strip(), f"item {entry} COUNT mismatch"


# ---------------------------------------------------------------------------
# 6. Exactly one `,deleted` item row, appended after kept block in ascending
#    ENTRY order, with blank measure/fit and SCORE/COUNT filled.
# ---------------------------------------------------------------------------

def test_exactly_one_deleted_row_last_blank_measure(tmp_path):
    _run(tmp_path, idfile_path=DEMO_IDFILE)
    rows = _item_rows(tmp_path)

    deleted = [r for r in rows if r[15].strip() == "deleted"]
    kept = [r for r in rows if r[15].strip() == "kept"]
    assert len(deleted) == 1
    assert len(kept) == 5

    # deleted row is after all kept rows
    last_kept_idx = max(rows.index(r) for r in kept)
    del_idx = rows.index(deleted[0])
    assert del_idx > last_kept_idx

    # ENTRY ascending within kept + deleted combined (entries 1-6)
    entries = [int(r[0].strip()) for r in rows]
    assert entries == sorted(entries)

    d = deleted[0]
    assert d[0].strip() == "6"           # ENTRY
    assert d[1].strip() != ""            # SCORE filled
    assert d[2].strip() != ""            # COUNT filled
    assert d[3].strip() == ""            # MEASURE blank
    assert d[4].strip() == ""            # S.E. blank
    assert d[5].strip() == ""            # INFIT MNSQ blank
    assert d[6].strip() == ""            # INFIT ZSTD blank


# ---------------------------------------------------------------------------
# 7. Summary: COUNTS,ITEM DELETED=1; console prints "Deleted items" line;
#    summary ITEM COUNT = NI - 1 = 5
# ---------------------------------------------------------------------------

def test_summary_item_deleted_count(tmp_path):
    out, err, code = _run(tmp_path, idfile_path=DEMO_IDFILE)
    assert code == 0, err

    s = _summary_dict(tmp_path)
    assert s[("COUNTS", "ITEM DELETED")] == "1"
    assert s[("ITEM", "COUNT")] == "5"  # NI - 1

    assert "Deleted items      : 1" in out


# ---------------------------------------------------------------------------
# 8. option_table_15.3.csv contains no row for the deleted item
# ---------------------------------------------------------------------------

def test_option_table_no_deleted_item(tmp_path):
    _run(tmp_path, idfile_path=DEMO_IDFILE)
    _, rows = _read_csv(tmp_path / "option_table_15.3.csv")
    item_entries = [r[0].strip() for r in rows if r]
    assert "6" not in item_entries


# ---------------------------------------------------------------------------
# 9. Deleted item absent from wright_map_measure.csv
# ---------------------------------------------------------------------------

def test_wright_map_no_deleted_item(tmp_path):
    _run(tmp_path, idfile_path=DEMO_IDFILE)
    _, rows = _read_csv(tmp_path / "wright_map_measure.csv")
    # ITEM_ENTRIES column (index 7 in MEASURE_HEADER_ROW_1) lists item entries
    item_entry_col = 7  # ITEM_ENTRIES is the last column
    for r in rows:
        if len(r) > item_entry_col:
            assert r[item_entry_col].strip() != "6", f"Item 6 found in wright_map_measure: {r}"


# ---------------------------------------------------------------------------
# 10. Workbook: 15.1 sheet has STATUS last; deleted row present with None
#     measure cells; auto_filter.ref starts at "A2"
# ---------------------------------------------------------------------------

def test_workbook_15_1_sheet(tmp_path):
    _run(tmp_path, idfile_path=DEMO_IDFILE)
    wb = openpyxl.load_workbook(tmp_path / "analysis_report.xlsx")
    ws = wb["15.1"]

    # STATUS is the last column (row 2, which is the second header row)
    header2 = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    assert header2[-1] == "STATUS"

    # auto_filter starts at A2
    assert ws.auto_filter.ref is not None
    assert ws.auto_filter.ref.startswith("A2")

    # Find the deleted row (STATUS=="deleted")
    del_row = None
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[-1] == "deleted":
            del_row = row
            break
    assert del_row is not None, "No deleted row in workbook 15.1 sheet"

    # MEASURE cell (col 4, index 3) must be None (empty)
    assert del_row[3] is None, f"Deleted row MEASURE should be None, got {del_row[3]}"
    # ENTRY filled, SCORE filled, COUNT filled
    assert del_row[0] == 6
    assert del_row[1] is not None
    assert del_row[2] is not None


# ---------------------------------------------------------------------------
# 11. Idfile that deletes every item -> exit 2
# ---------------------------------------------------------------------------

def test_idfile_all_items_deleted_exits_2(tmp_path):
    idf = tmp_path / "all.TXT"
    idf.write_text("1\n2\n3\n4\n5\n6\n", encoding="utf-8")
    out, err, code = _run(tmp_path / "out", idfile_path=str(idf))
    assert code == 2
    assert "no items remain" in err
