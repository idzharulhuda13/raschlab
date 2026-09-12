import csv
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

from raschlab.distractor import option_table


# Number formats for numeric table columns, keyed by the second header-row name
# used on the item/person ('15.1', 'person') and option ('15.3') sheets.
NUMBER_FORMATS = {
    "NUMBER": "0",       # ENTRY NUMBER
    "SCORE": "0",        # TOTAL SCORE
    "COUNT": "0",        # TOTAL COUNT / option data count
    "VALUE": "0",        # option score value
    "%": "0",            # option count percentage
    "MEASURE": "0.00",   # JMLE MEASURE
    "S.E.": "0.00",      # MODEL S.E.
    "MNSQ": "0.00",      # INFIT/OUTFIT MNSQ
    "ZSTD": "0.00",      # INFIT/OUTFIT ZSTD
    "CORR.": "0.00",     # PTMEASUR-AL CORR.
    "EXP.": "0.00",      # PTMEASUR-AL EXP.
    "OBS%": "0.0",       # EXACT MATCH OBS%
    "EXP%": "0.0",       # EXACT MATCH EXP%
    "ABILITY MEAN": "0.00",
    "ABILITY PSD": "0.00",
    "SE MEAN": "0.00",
    "INFT MNSQ": "0.00",
    "OUTF MNSQ": "0.00",
    "PTMA CORR": "0.00",
}


def number_format_for_header(header):
    """Number format for a column given its second header-row label, or None."""
    if header is None:
        return None
    return NUMBER_FORMATS.get(str(header).strip().upper())


def coerce_cell(value, num_fmt):
    """Return (cell_value, number_format) turning a display string into a number.

    The row builders emit already-rounded display strings (and the CSV output
    must keep them), so numeric columns are converted back to real numbers here
    when the workbook is written. Non-numeric text is passed through untouched
    and genuinely empty cells stay empty.
    """
    if value is None or value == "":
        return None, num_fmt
    if isinstance(value, bool):
        return value, None
    if isinstance(value, (int, float, np.integer, np.floating)):
        num = float(value)
    elif isinstance(value, str):
        try:
            num = float(value.strip())
        except ValueError:
            return value, None
    else:
        return value, None
    if np.isnan(num) or np.isinf(num):
        return "", None
    if num_fmt == "0":
        return int(round(num)), "0"
    return float(num), num_fmt


def summary_value_format(value):
    """Number format for a numeric summary VALUE cell (int -> '0', else '0.00')."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, np.integer)):
        return "0"
    if isinstance(value, (float, np.floating)):
        return "0.00"
    if isinstance(value, str):
        text = value.strip()
        try:
            float(text)
        except ValueError:
            return None
        return "0.00" if "." in text else "0"
    return None


def apply_column_widths(ws, header_rows, data_rows):
    """Set column widths wide enough for the header text and the data."""
    widths = {}
    for row in list(header_rows) + [list(r) for r in data_rows]:
        for i, val in enumerate(row, start=1):
            if val is None or val == "":
                continue
            widths[i] = max(widths.get(i, 0), len(str(val)))
    for i, width in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = width + 2


def fill_sheet(ws, header_rows, data_rows, fmt_for, freeze_panes="A3"):
    """Write header rows plus data rows, applying number formats and widths."""
    for header_row in header_rows:
        ws.append(list(header_row))
    for row in data_rows:
        ws.append(list(row))

    names = header_rows[-1] if header_rows else []
    first_data_row = len(header_rows) + 1
    for r in range(first_data_row, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            name = names[c - 1] if c - 1 < len(names) else None
            cell = ws.cell(row=r, column=c)
            fmt = fmt_for(name, cell.value)
            if fmt is None:
                continue
            cell.value, cell.number_format = coerce_cell(cell.value, fmt)

    ws.freeze_panes = freeze_panes
    apply_column_widths(ws, header_rows, data_rows)


# Two-row headers matching Winsteps / team sheets
ITEM_HEADER_ROW_1 = [
    "ENTRY", "TOTAL", "TOTAL", "JMLE", "MODEL", "INFIT", "", "OUTFIT", "", "PTMEASUR-AL", "", "EXACT", "MATCH"
]
ITEM_HEADER_ROW_2 = [
    "NUMBER", "SCORE", "COUNT", "MEASURE", "S.E.", "MNSQ", "ZSTD", "MNSQ", "ZSTD", "CORR.", "EXP.", "OBS%", "EXP%"
]

PERSON_HEADER_ROW_1 = [
    "ENTRY", "TOTAL", "TOTAL", "JMLE", "MODEL", "INFIT", "", "OUTFIT", "", "PTMEASUR-AL", "", "EXACT", "MATCH", ""
]
PERSON_HEADER_ROW_2 = [
    "NUMBER", "SCORE", "COUNT", "MEASURE", "S.E.", "MNSQ", "ZSTD", "MNSQ", "ZSTD", "CORR.", "EXP.", "OBS%", "EXP%", "PERSON", "RANK"
]

OPTION_HEADER_ROW_1 = [
    "ENTRY", "DATA", "SCORE", "DATA", "", "ABILITY", "", "S.E.", "INFT", "OUTF", "PTMA", ""
]
OPTION_HEADER_ROW_2 = [
    "NUMBER", "CODE", "VALUE", "COUNT", "%", "ABILITY MEAN", "ABILITY PSD", "SE MEAN", "INFT MNSQ", "OUTF MNSQ", "PTMA CORR", "ITEM"
]


class TableRow(dict):
    """Row dictionary that preserves key ordering and supports standard alias lookups."""
    _ALIASES = {
        "NUMBER": "ENTRY",
        "ENTRY NUMBER": "ENTRY",
        "ITEM": "ENTRY",
        "TOTAL SCORE": "SCORE",
        "TOTAL COUNT": "COUNT",
        "JMLE MEASURE": "MEASURE",
        "MODEL S.E.": "S.E.",
        "SE": "S.E.",
        "MODEL SE": "S.E.",
        "PTMEASUR-AL CORR.": "CORR.",
        "PTMA CORR": "CORR.",
        "PTMEASUR-AL EXP.": "EXP.",
        "EXACT MATCH OBS%": "OBS%",
        "EXACT MATCH EXP%": "EXP%",
        "DATA COUNT": "COUNT",
        "DATA_COUNT": "COUNT",
        "DATA%": "%",
        "DATA_PCT": "%",
    }

    def __getitem__(self, key):
        if super().__contains__(key):
            return super().__getitem__(key)
        canon = self._ALIASES.get(key)
        if canon and super().__contains__(canon):
            return super().__getitem__(canon)
        raise KeyError(key)

    def get(self, key, default=None):
        try:
            return self[key]
        except KeyError:
            return default

    def __contains__(self, key):
        return super().__contains__(key) or (key in self._ALIASES and super().__contains__(self._ALIASES[key]))


class TableRowList(list):
    """List of TableRows with optional metadata attributes."""
    def __init__(self, items=(), n_extreme_excluded=0):
        super().__init__(items)
        self.n_extreme_excluded = n_extreme_excluded
        self.n_extreme = n_extreme_excluded


def _fmt(val, decimals):
    try:
        f = float(val)
        if np.isnan(f) or np.isinf(f):
            return ""
        rounded = round(f, decimals)
        if rounded == 0.0:
            rounded = 0.0
        return f"{rounded:.{decimals}f}"
    except (ValueError, TypeError):
        return str(val) if val is not None else ""


def item_table_rows(
    X,
    mask,
    key,
    item_measures,
    fit_item,
    keep=None,
    person_measures=None,
    extra_cols=None,
    digits=2,
):
    """Generate item table rows matching Winsteps Table 15.1.

    Parameters
    ----------
    X : array-like of shape (P, I)
        Scored response matrix (0/1 or NaN).
    mask : array-like of shape (P, I)
        Boolean validity mask.
    key : str or sequence of length I
        Scoring key.
    item_measures : array-like of shape (I,)
        Item calibration measures (logits).
    fit_item : dict
        Item fit statistics dict with keys 'se', 'infit_mnsq', 'infit_zstd',
        'outfit_mnsq', 'outfit_zstd'.
    keep : array-like of shape (P,), optional
        Boolean calibration mask for persons.
    person_measures : array-like of shape (P,), optional
        Person ability measures for computing point-measure correlation.
        If None, Cohen's PROX starting values are used.
    extra_cols : dict of str -> list, optional
        Extra columns (e.g. {'KET': [...]}) to append to each row.
    digits : int, optional
        Decimal digits for MEASURE and S.E. (default: 2).

    Returns
    -------
    list of TableRow
        List of dicts with 13 keys corresponding to Table 15.1.
    """
    X = np.asarray(X, dtype=float)
    mask = np.asarray(mask, dtype=bool)
    d = np.asarray(item_measures, dtype=float)
    P, I = X.shape

    if keep is not None:
        kp = np.asarray(keep, dtype=bool)
    else:
        kp = np.ones(P, dtype=bool)

    # Derive person measures if not supplied
    if person_measures is not None:
        b = np.asarray(person_measures, dtype=float)
    elif "person_measures" in fit_item:
        b = np.asarray(fit_item["person_measures"], dtype=float)
    else:
        m_i = np.sum(mask, axis=1)
        s_i = np.sum(np.where(mask, X, 0.0), axis=1)
        ratio_i = np.clip((s_i + 0.5) / np.maximum(m_i - s_i + 0.5, 1e-12), 1e-12, 1e12)
        b = np.log(ratio_i)

    rows = []
    for j in range(I):
        resp_mask = kp & mask[:, j]
        score_val = int(np.sum(np.where(resp_mask, X[:, j], 0.0)))
        count_val = int(np.sum(resp_mask))

        meas_val = float(d[j])
        se_val = float(fit_item["se"][j])
        infit_mnsq = float(fit_item["infit_mnsq"][j])
        infit_zstd = float(fit_item["infit_zstd"][j])
        outfit_mnsq = float(fit_item["outfit_mnsq"][j])
        outfit_zstd = float(fit_item["outfit_zstd"][j])

        # Point-measure correlation over persons who answered item j in calibration set
        b_resp = b[resp_mask]
        x_resp = X[resp_mask, j]
        if fit_item is not None and "ptmeas" in fit_item and fit_item["ptmeas"] is not None:
            corr_val = float(fit_item["ptmeas"][j])
        elif len(x_resp) > 1 and np.std(x_resp, ddof=0) > 1e-12 and np.std(b_resp, ddof=0) > 1e-12:
            corr_val = float(np.corrcoef(x_resp, b_resp)[0, 1])
        else:
            corr_val = 0.0

        # Expected point-measure correlation (EXP.)
        if fit_item is not None and "exp" in fit_item and fit_item["exp"] is not None:
            exp_val = float(fit_item["exp"][j])
        elif len(b_resp) > 0:
            diff = b_resp - d[j]
            P_j = 1.0 / (1.0 + np.exp(-np.clip(diff, -30.0, 30.0)))
            b_bar = np.mean(b_resp)
            P_bar = np.mean(P_j)
            num = np.mean((b_resp - b_bar) * (P_j - P_bar))
            conv = np.sqrt(P_bar * (1.0 - P_bar))
            sd_b = np.std(b_resp, ddof=0)
            denom = sd_b * conv
            exp_val = float(num / denom) if denom > 1e-12 else 0.0
        else:
            exp_val = 0.0
        exp_corr = _fmt(exp_val, 2)

        # Exact match percentages: OBS% and EXP%
        if fit_item is not None and "obs_pct" in fit_item and fit_item["obs_pct"] is not None:
            obs_pct = float(fit_item["obs_pct"][j])
        elif len(b_resp) > 0:
            diff = b_resp - d[j]
            P_j = 1.0 / (1.0 + np.exp(-np.clip(diff, -30.0, 30.0)))
            expected_resp = (P_j >= 0.5).astype(float)
            obs_match = (x_resp == expected_resp)
            obs_pct = float(np.mean(obs_match) * 100)
        else:
            obs_pct = 0.0

        if fit_item is not None and "exp_pct" in fit_item and fit_item["exp_pct"] is not None:
            exp_pct = float(fit_item["exp_pct"][j])
        elif len(b_resp) > 0:
            diff = b_resp - d[j]
            P_j = 1.0 / (1.0 + np.exp(-np.clip(diff, -30.0, 30.0)))
            exp_pct = float(np.mean(np.maximum(P_j, 1.0 - P_j)) * 100)
        else:
            exp_pct = 0.0

        row = TableRow()
        row["ENTRY"] = j + 1
        row["SCORE"] = score_val
        row["COUNT"] = count_val
        row["MEASURE"] = _fmt(meas_val, digits)
        row["S.E."] = _fmt(se_val, digits)
        row["INFIT MNSQ"] = _fmt(infit_mnsq, 2)
        row["INFIT ZSTD"] = _fmt(infit_zstd, 2)
        row["OUTFIT MNSQ"] = _fmt(outfit_mnsq, 2)
        row["OUTFIT ZSTD"] = _fmt(outfit_zstd, 2)
        row["CORR."] = _fmt(corr_val, 2)
        row["EXP."] = exp_corr
        row["OBS%"] = _fmt(obs_pct, 1)
        row["EXP%"] = _fmt(exp_pct, 1)

        if extra_cols:
            for k, vals in extra_cols.items():
                row[k] = vals[j] if j < len(vals) else ""

        rows.append(row)

    return rows


def person_table_rows(
    X,
    mask,
    key,
    person_measures,
    fit_person,
    keep,
    labels,
    item_measures=None,
    extra_cols=None,
    digits=2,
    order=None,
    rank_letters=False,
):
    """Generate person table rows.
    Extreme persons are excluded and counted separately.

    Parameters
    ----------
    X : array-like of shape (P, I)
        Scored response matrix.
    mask : array-like of shape (P, I)
        Boolean validity mask.
    key : str or sequence
        Scoring key.
    person_measures : array-like of shape (P,)
        Person ability measures (logits).
    fit_person : dict
        Person fit statistics dict with keys 'se', 'infit_mnsq', 'infit_zstd',
        'outfit_mnsq', 'outfit_zstd'.
    keep : array-like of shape (P,)
        Boolean calibration mask.
    labels : list of str of length P
        Person labels.
    item_measures : array-like of shape (I,), optional
        Item difficulty measures for computing CORR., OBS%, EXP%.
    extra_cols : dict of str -> list, optional
        Extra columns to append to each row.
    digits : int, optional
        Decimal digits for MEASURE and S.E. (default: 2).
    order : sequence of int, optional
        Sequence of person indices defining the output row order.
    rank_letters : bool, optional
        Whether to append RANK column with rank letters (default: False).

    Returns
    -------
    TableRowList of TableRow
        Rows for non-extreme kept persons, with attribute n_extreme_excluded.
    """
    X = np.asarray(X, dtype=float)
    mask = np.asarray(mask, dtype=bool)
    pm = np.asarray(person_measures, dtype=float)
    P, I = X.shape

    scores = np.nansum(X, axis=1)
    counts = np.sum(mask, axis=1)
    is_extreme = (counts == 0) | (scores == 0) | (scores == counts)

    if keep is not None:
        kp = np.asarray(keep, dtype=bool)
        valid = kp & ~is_extreme
        n_extreme = int(np.sum(kp & is_extreme))
    else:
        valid = ~is_extreme
        n_extreme = int(np.sum(is_extreme))

    valid_indices = np.where(valid)[0]

    d = np.asarray(item_measures, dtype=float) if item_measures is not None else None

    rows = TableRowList(n_extreme_excluded=n_extreme)
    for p_idx, i in enumerate(valid_indices):
        entry = i + 1
        score_val = int(scores[i])
        count_val = int(counts[i])
        meas_val = float(pm[i])
        se_val = float(fit_person["se"][p_idx])
        infit_mnsq = float(fit_person["infit_mnsq"][p_idx])
        infit_zstd = float(fit_person["infit_zstd"][p_idx])
        outfit_mnsq = float(fit_person["outfit_mnsq"][p_idx])
        outfit_zstd = float(fit_person["outfit_zstd"][p_idx])

        # Correlation and match percentages across answered items
        if d is not None:
            items_i = np.where(mask[i])[0]
            x_i = X[i, items_i]
            d_i = d[items_i]
            if len(x_i) > 1 and np.std(x_i, ddof=0) > 1e-12 and np.std(d_i, ddof=0) > 1e-12:
                # In Winsteps, person PTMEASUR-AL CORR is correlation with item easiness (-d)
                corr_val = float(np.corrcoef(x_i, -d_i)[0, 1])
            else:
                corr_val = 0.0

            # Expected point-measure correlation (EXP.)
            if fit_person is not None and "exp" in fit_person and fit_person["exp"] is not None:
                exp_val = float(fit_person["exp"][p_idx])
            elif len(d_i) > 0:
                diff = pm[i] - d_i
                P_i = 1.0 / (1.0 + np.exp(-np.clip(diff, -30.0, 30.0)))
                d_bar = np.mean(d_i)
                P_bar = np.mean(P_i)
                num = np.mean((d_i - d_bar) * (P_i - P_bar))
                conv = np.sqrt(P_bar * (1.0 - P_bar))
                sd_d = np.std(d_i, ddof=0)
                denom = sd_d * conv
                exp_val = float(-num / denom) if denom > 1e-12 else 0.0
            else:
                exp_val = 0.0
            exp_corr = _fmt(exp_val, 2)

            diff = pm[i] - d_i
            P_i = 1.0 / (1.0 + np.exp(-np.clip(diff, -30.0, 30.0)))
            expected_resp = (P_i >= 0.5).astype(float)
            obs_match = (x_i == expected_resp)
            obs_pct = float(np.mean(obs_match) * 100)
            exp_pct = float(np.mean(np.maximum(P_i, 1.0 - P_i)) * 100)
        else:
            corr_val = 0.0
            exp_corr = _fmt(0.0, 2)
            obs_pct = 0.0
            exp_pct = 0.0

        label_val = str(labels[i]) if labels is not None and i < len(labels) else str(entry)

        row = TableRow()
        row["ENTRY"] = entry
        row["SCORE"] = score_val
        row["COUNT"] = count_val
        row["MEASURE"] = _fmt(meas_val, digits)
        row["S.E."] = _fmt(se_val, digits)
        row["INFIT MNSQ"] = _fmt(infit_mnsq, 2)
        row["INFIT ZSTD"] = _fmt(infit_zstd, 2)
        row["OUTFIT MNSQ"] = _fmt(outfit_mnsq, 2)
        row["OUTFIT ZSTD"] = _fmt(outfit_zstd, 2)
        row["CORR."] = _fmt(corr_val, 2)
        row["EXP."] = exp_corr
        row["OBS%"] = _fmt(obs_pct, 1)
        row["EXP%"] = _fmt(exp_pct, 1)
        row["PERSON"] = label_val

        if extra_cols:
            for k, vals in extra_cols.items():
                row[k] = vals[i] if i < len(vals) else ""

        rows.append(row)

    if order is not None:
        row_map = {i: r for i, r in zip(valid_indices, rows)}
        rows = TableRowList(
            [row_map[idx] for idx in order if idx in row_map],
            n_extreme_excluded=rows.n_extreme_excluded,
        )

    if rank_letters:
        n = len(rows)
        for idx, row in enumerate(rows):
            if idx < 26:
                row["RANK"] = chr(ord("A") + idx)
            elif idx >= n - 26:
                row["RANK"] = chr(ord("a") + (n - 1 - idx))
            else:
                row["RANK"] = ""

    return rows


def option_rows(*args, item_labels=None, **kwargs):
    """Wrap distractor.option_table output, renaming keys to team's labels:
    NUMBER, CODE, VALUE, COUNT, %, ABILITY MEAN, ABILITY PSD, SE MEAN,
    INFT MNSQ, OUTF MNSQ, PTMA CORR, ITEM.
    """
    if len(args) == 1 and isinstance(args[0], list) and (len(args[0]) == 0 or isinstance(args[0][0], dict)):
        raw_rows = args[0]
        if item_labels is not None:
            for r in raw_rows:
                num = r.get("NUMBER")
                if num is not None:
                    try:
                        j = int(num) - 1
                        if 0 <= j < len(item_labels):
                            r["ITEM"] = str(item_labels[j])
                    except (ValueError, TypeError):
                        pass
    else:
        raw_rows = option_table(*args, item_labels=item_labels, **kwargs)

    key_map = {
        "NUMBER": "NUMBER",
        "CODE": "CODE",
        "VALUE": "VALUE",
        "DATA_COUNT": "COUNT",
        "DATA COUNT": "COUNT",
        "COUNT": "COUNT",
        "DATA_PCT": "%",
        "DATA%": "%",
        "%": "%",
        "ABILITY_MEAN": "ABILITY MEAN",
        "ABILITY MEAN": "ABILITY MEAN",
        "ABILITY_PSD": "ABILITY PSD",
        "ABILITY PSD": "ABILITY PSD",
        "SE_MEAN": "SE MEAN",
        "SE MEAN": "SE MEAN",
        "INFT_MNSQ": "INFT MNSQ",
        "INFT MNSQ": "INFT MNSQ",
        "OUTF_MNSQ": "OUTF MNSQ",
        "OUTF MNSQ": "OUTF MNSQ",
        "PTMA_CORR": "PTMA CORR",
        "PTMA CORR": "PTMA CORR",
        "ITEM": "ITEM",
    }

    two_dec_keys = {
        "ABILITY MEAN",
        "ABILITY PSD",
        "SE MEAN",
        "INFT MNSQ",
        "OUTF MNSQ",
        "PTMA CORR",
    }
    int_keys = {
        "NUMBER",
        "VALUE",
        "COUNT",
        "%",
        "DATA COUNT",
        "DATA%",
    }

    out = []
    for r in raw_rows:
        renamed = TableRow()
        for k, v in r.items():
            target_key = key_map.get(k, k)
            if target_key in two_dec_keys:
                renamed[target_key] = _fmt(v, 2)
            elif target_key in int_keys:
                if v is None or v == "":
                    renamed[target_key] = ""
                else:
                    try:
                        renamed[target_key] = int(round(float(v)))
                    except (ValueError, TypeError):
                        renamed[target_key] = v
            else:
                renamed[target_key] = v
        out.append(renamed)
    return out


def summary_rows(item_summary, person_summary, counts_info=None):
    """Build list of (section, label, value) triples from item and person summaries.

    Parameters
    ----------
    item_summary : dict
        Output from raschlab.summary.item_summary.
    person_summary : dict
        Output from raschlab.summary.person_summary.
    counts_info : dict, optional
        Counts for deleted, lacking, extreme_min, extreme_max.

    Returns
    -------
    list of tuple (str, str, Any)
    """
    rows = []

    # Item section
    im = item_summary.get("measure", {})
    ise = item_summary.get("se", {})
    i_inf = item_summary.get("infit_mnsq", {})
    i_outf = item_summary.get("outfit_mnsq", {})
    i_real = item_summary.get("real", {})
    i_mod = item_summary.get("model", {})

    rows.append(("ITEM", "COUNT", item_summary.get("count", 0)))
    rows.append(("ITEM MEASURE", "MEAN", round(float(im.get("mean", 0.0)), 2)))
    rows.append(("ITEM MEASURE", "SEM", round(float(im.get("sem", 0.0)), 2)))
    rows.append(("ITEM MEASURE", "P.SD", round(float(im.get("psd", 0.0)), 2)))
    rows.append(("ITEM MEASURE", "MAX", round(float(im.get("max", 0.0)), 2)))
    rows.append(("ITEM MEASURE", "MIN", round(float(im.get("min", 0.0)), 2)))

    rows.append(("ITEM MODEL S.E.", "MEAN", round(float(ise.get("mean", 0.0)), 2)))
    rows.append(("ITEM MODEL S.E.", "SEM", round(float(ise.get("sem", 0.0)), 2)))
    rows.append(("ITEM MODEL S.E.", "P.SD", round(float(ise.get("psd", 0.0)), 2)))
    rows.append(("ITEM MODEL S.E.", "MAX", round(float(ise.get("max", 0.0)), 2)))
    rows.append(("ITEM MODEL S.E.", "MIN", round(float(ise.get("min", 0.0)), 2)))

    rows.append(("ITEM INFIT MNSQ", "MEAN", round(float(i_inf.get("mean", 0.0)), 2)))
    rows.append(("ITEM INFIT MNSQ", "SD", round(float(i_inf.get("sd", 0.0)), 2)))
    rows.append(("ITEM OUTFIT MNSQ", "MEAN", round(float(i_outf.get("mean", 0.0)), 2)))
    rows.append(("ITEM OUTFIT MNSQ", "SD", round(float(i_outf.get("sd", 0.0)), 2)))

    rows.append(("ITEM REAL", "RMSE", round(float(i_real.get("rmse", 0.0)), 2)))
    rows.append(("ITEM REAL", "TRUE SD", round(float(i_real.get("true_sd", 0.0)), 2)))
    rows.append(("ITEM REAL", "SEPARATION", round(float(i_real.get("separation", 0.0)), 2)))
    rows.append(("ITEM REAL", "RELIABILITY", round(float(i_real.get("reliability", 0.0)), 2)))

    rows.append(("ITEM MODEL", "RMSE", round(float(i_mod.get("rmse", 0.0)), 2)))
    rows.append(("ITEM MODEL", "TRUE SD", round(float(i_mod.get("true_sd", 0.0)), 2)))
    rows.append(("ITEM MODEL", "SEPARATION", round(float(i_mod.get("separation", 0.0)), 2)))
    rows.append(("ITEM MODEL", "RELIABILITY", round(float(i_mod.get("reliability", 0.0)), 2)))

    # Person section
    pm = person_summary.get("measure", {})
    pse = person_summary.get("se", {})
    p_inf = person_summary.get("infit_mnsq", {})
    p_outf = person_summary.get("outfit_mnsq", {})
    p_real = person_summary.get("real", {})
    p_mod = person_summary.get("model", {})

    rows.append(("PERSON", "COUNT", person_summary.get("count", 0)))
    rows.append(("PERSON MEASURE", "MEAN", round(float(pm.get("mean", 0.0)), 2)))
    rows.append(("PERSON MEASURE", "SEM", round(float(pm.get("sem", 0.0)), 2)))
    rows.append(("PERSON MEASURE", "P.SD", round(float(pm.get("psd", 0.0)), 2)))
    rows.append(("PERSON MEASURE", "MAX", round(float(pm.get("max", 0.0)), 2)))
    rows.append(("PERSON MEASURE", "MIN", round(float(pm.get("min", 0.0)), 2)))

    rows.append(("PERSON MODEL S.E.", "MEAN", round(float(pse.get("mean", 0.0)), 2)))
    rows.append(("PERSON MODEL S.E.", "SEM", round(float(pse.get("sem", 0.0)), 2)))
    rows.append(("PERSON MODEL S.E.", "P.SD", round(float(pse.get("psd", 0.0)), 2)))
    rows.append(("PERSON MODEL S.E.", "MAX", round(float(pse.get("max", 0.0)), 2)))
    rows.append(("PERSON MODEL S.E.", "MIN", round(float(pse.get("min", 0.0)), 2)))

    rows.append(("PERSON INFIT MNSQ", "MEAN", round(float(p_inf.get("mean", 0.0)), 2)))
    rows.append(("PERSON INFIT MNSQ", "SD", round(float(p_inf.get("sd", 0.0)), 2)))
    rows.append(("PERSON OUTFIT MNSQ", "MEAN", round(float(p_outf.get("mean", 0.0)), 2)))
    rows.append(("PERSON OUTFIT MNSQ", "SD", round(float(p_outf.get("sd", 0.0)), 2)))

    rows.append(("PERSON REAL", "RMSE", round(float(p_real.get("rmse", 0.0)), 2)))
    rows.append(("PERSON REAL", "TRUE SD", round(float(p_real.get("true_sd", 0.0)), 2)))
    rows.append(("PERSON REAL", "SEPARATION", round(float(p_real.get("separation", 0.0)), 2)))
    rows.append(("PERSON REAL", "RELIABILITY", round(float(p_real.get("reliability", 0.0)), 2)))

    rows.append(("PERSON MODEL", "RMSE", round(float(p_mod.get("rmse", 0.0)), 2)))
    rows.append(("PERSON MODEL", "TRUE SD", round(float(p_mod.get("true_sd", 0.0)), 2)))
    rows.append(("PERSON MODEL", "SEPARATION", round(float(p_mod.get("separation", 0.0)), 2)))
    rows.append(("PERSON MODEL", "RELIABILITY", round(float(p_mod.get("reliability", 0.0)), 2)))

    if "raw_score_corr" in person_summary and person_summary["raw_score_corr"] is not None:
        rows.append(("PERSON CORR", "RAW SCORE TO MEASURE CORRELATION", round(float(person_summary["raw_score_corr"]), 2)))

    # Counts section
    extreme_count = person_summary.get("n_extreme_excluded", 0)
    rows.append(("COUNTS", "EXTREME EXCLUDED", extreme_count))

    if counts_info:
        for k, v in counts_info.items():
            rows.append(("COUNTS", k.upper(), v))
    else:
        for k in ("lacking", "deleted", "extreme_min", "extreme_max"):
            if k in person_summary:
                rows.append(("COUNTS", k.upper(), person_summary[k]))

    return rows


def write_csv(rows, path, header_rows=None):
    """Write rows to a CSV file in dict key order.

    Parameters
    ----------
    rows : list of dict or list of list/tuple
        Data rows.
    path : str
        Target file path.
    header_rows : list of list of str, optional
        One or two header rows to write at the beginning.
    """
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if header_rows:
            for hr in header_rows:
                writer.writerow(hr)
        elif rows and isinstance(rows[0], dict):
            writer.writerow(list(rows[0].keys()))

        for r in rows:
            if isinstance(r, dict):
                writer.writerow(list(r.values()))
            elif isinstance(r, (list, tuple)):
                writer.writerow(r)


def write_workbook(path, item_rows, person_rows, option_rows, summary_rows):
    """Write XLSX workbook containing sheets: '15.1', 'person', '15.3', 'summary'.

    Each sheet starts with its own two-row header where applicable. Header rows
    are frozen, numeric columns are written as real numbers carrying an explicit
    number format (so decimals survive a paste into Google Sheets) and column
    widths fit the header text. Text labels and genuinely empty cells are kept
    as-is.
    """
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    def table_fmt(name, value):
        return number_format_for_header(name)

    def summary_fmt(name, value):
        return summary_value_format(value)

    # Sheet 15.1 (Items)
    ws_item = wb.create_sheet(title="15.1")
    fill_sheet(
        ws_item,
        [ITEM_HEADER_ROW_1, ITEM_HEADER_ROW_2],
        [list(r.values()) for r in item_rows],
        table_fmt,
    )

    # Sheet person (Persons)
    ws_person = wb.create_sheet(title="person")
    fill_sheet(
        ws_person,
        [PERSON_HEADER_ROW_1, PERSON_HEADER_ROW_2],
        [list(r.values()) for r in person_rows],
        table_fmt,
    )

    # Sheet 15.3 (Options)
    ws_opt = wb.create_sheet(title="15.3")
    fill_sheet(
        ws_opt,
        [OPTION_HEADER_ROW_1, OPTION_HEADER_ROW_2],
        [list(r.values()) for r in option_rows],
        table_fmt,
    )

    # Sheet summary (two header rows: labels then a blank spacer row)
    ws_sum = wb.create_sheet(title="summary")
    fill_sheet(
        ws_sum,
        [["SECTION", "STATISTIC", "VALUE"], ["", "", ""]],
        [list(trip) for trip in summary_rows],
        summary_fmt,
    )

    wb.remove(default_sheet)
    wb.save(path)
